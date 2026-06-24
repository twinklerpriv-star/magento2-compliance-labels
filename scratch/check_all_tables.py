import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

def find_all_tables(file_path):
    if not os.path.exists(file_path):
        print(f"Error: '{file_path}' does not exist.")
        return
        
    wb = openpyxl.load_workbook(file_path)
    for sheet_name in wb.sheetnames:
        print(f"\nSheet: '{sheet_name}'")
        ws = wb[sheet_name]
        print(f"Tables count: {len(ws.tables)}")
        for name, table_ref in ws.tables.items():
            print(f"  Table Name: {name}")
            print(f"  Range: {table_ref}")
            # Get table object
            table = ws.tables[name]
            if hasattr(table, 'tableStyleInfo') and table.tableStyleInfo:
                print(f"  Style Name: {table.tableStyleInfo.name}")
                print(f"  Show Row Stripes: {table.tableStyleInfo.showRowStripes}")
                print(f"  Show Column Stripes: {table.tableStyleInfo.showColumnStripes}")

def main():
    file_path = r"C:\Users\thomas.winkler\Desktop\Projekte\Google Antigravity\ELEKTROPEPI\GEWAEHRLEISTUNG_GARANTIELABEL\alle gelisteten Artikel - Stand 24.06.2026.xlsx"
    find_all_tables(file_path)

if __name__ == "__main__":
    main()
